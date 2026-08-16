#!/usr/bin/env python3
"""Rent chart data workbook: HD mix-weighted series + UW paths + L5 new leases."""
import openpyxl, json, datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
wb=openpyxl.Workbook()
F=lambda **k: Font(name="Arial", **{"size":10, **k})
TITLE=F(bold=True,size=13); H1=F(bold=True,color='FFFFFF'); BOLD=F(bold=True)
BLUE=F(color='0000FF'); BLACK=F(); GREEN=F(color='008000'); NOTE=F(italic=True,size=9,color='808080')
HDR=PatternFill('solid',fgColor='1F3864'); SUB=PatternFill('solid',fgColor='D9E2F2')
M0='$#,##0'; M2='$#,##0.00'; D='mm/dd/yyyy'; NUM='#,##0'
def sheet(n,w):
    ws=wb.create_sheet(n)
    for i,x in enumerate(w,1): ws.column_dimensions[get_column_letter(i)].width=x
    ws.sheet_view.showGridLines=False
    return ws
def put(ws,cell,v,f=BLACK,fmt=None,fill=None,wrap=False):
    c=ws[cell]; c.value=v; c.font=f
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    if wrap: c.alignment=Alignment(wrap_text=True,vertical='top')
def hdr(ws,row,cols,texts):
    for col,t in zip(cols,texts):
        c=ws[f'{col}{row}']; c.value=t; c.font=H1; c.fill=HDR
        c.alignment=Alignment(horizontal='center',wrap_text=True)

QLAB=["Q3 23","Q4 23","Q1 24","Q2 24","Q3 24","Q4 24","Q1 25","Q2 25","Q3 25","Q4 25","Q1 26","Q2 26","T90 (Jun-Aug 26)"]
SEA=[("S1_Seas",15,[1403,1403,1403,1403,1415,1351,1345,1509,1509,1325,1310,1544,1597]),
     ("A1_Seas",120,[1479,1479,1479,1610,1586,1432,1447,1536,1518,1394,1483,1660,1691]),
     ("A2_Seas",45,[1625,1625,1625,1625,1695,1676,1689,1771,1679,1556,1570,1692,1787]),
     ("B1_Seas",30,[1805,1805,1805,1805,1960,1765,1737,1910,1799,1726,1805,2030,2050]),
     ("B2_Seas",75,[1983,1983,1983,2045,2021,1945,2016,2007,1930,1838,1935,2125,2104]),
     ("B3a_Seas",15,[1934,1934,1934,1934,2173,2173,1898,1807,1945,2135,1978,1978,2075]),
     ("B3b_Seas",30,[1941,1941,1941,1941,2098,2098,1826,1958,1958,1866,1870,2102,2100]),
     ("C1a_Seas",10,[2371,2371,2371,2371,2355,2355,2132,2296,2186,2326,2344,2579,2579]),
     ("C1b_Seas",20,[2237,2237,2237,2237,2325,1855,2123,2248,2392,2280,2173,2554,2554])]
PQLAB=["Q3 23","Q4 23","Q1 24","Q2 24","Q3 24","Q4 24","Q1 25","Q2 25","Q3 25","Q4 25 QTD"]
PRE=[("1x1-pp (787sf)",106,[1493,1448,1434,1383,1373,1369,1342,1439,1465,1462]),
     ("2x2a-pp (1092sf)",100,[1725,1706,1685,1698,1802,1629,1682,1718,1811,1795]),
     ("2x2b-pp (1172sf)",30,[1895,1835,1798,1791,1818,1643,1698,1792,1838,1790]),
     ("3x2-pp (1291sf)",44,[2015,1867,1870,1872,2125,2030,2072,2079,2088,1970])]

# ---------------- HD_Weighting ----------------
ws=sheet('HD_Weighting',[2,20,9]+[10]*13+[3])
put(ws,'B2','HelloData Mix-Weighted Asking Rent — methodology = Seasons model Cash Flow (Annual) row 304',TITLE)
put(ws,'B3','Per floor plan HelloData asking by period (blue, from each model), weighted by rent-roll unit counts: Weighted Avg = SUMPRODUCT(values x units) / SUM(units). Identical convention to CF(Annual)!AP304 and Rent Analysis row 59.',NOTE,wrap=True)
put(ws,'B5','SEASONS AT MERIDIAN — source: TMG model Rent Analysis K9:W17 (HD pull 8/2026)',BOLD,fill=SUB)
hdr(ws,6,['B','C']+[get_column_letter(4+i) for i in range(13)],['Floor plan','Units']+QLAB)
r=7
for plan,u,vals in SEA:
    put(ws,f'B{r}',plan); put(ws,f'C{r}',u,BLUE,NUM)
    for i,v in enumerate(vals): put(ws,f'{get_column_letter(4+i)}{r}',v,BLUE,M0)
    r+=1
put(ws,f'B{r}','WEIGHTED AVG',BOLD)
put(ws,f'C{r}',f'=SUM(C7:C{r-1})',BOLD,NUM)
for i in range(13):
    col=get_column_letter(4+i)
    put(ws,f'{col}{r}',f'=SUMPRODUCT({col}7:{col}{r-1},$C$7:$C${r-1})/SUM($C$7:$C${r-1})',BOLD,M0)
SEA_W=r
put(ws,f'B{r+1}','Note: Q3 23–Q1 24 are rent-roll backfill (pre-delivery), not real HelloData — excluded from Chart_Data.',NOTE,wrap=True)
r+=3
put(ws,f'B{r}','PRELUDE AT PARAMOUNT — source: Prelude model Market Rent Summary K7:T11 (HD pull 11/2025)',BOLD,fill=SUB)
hdr(ws,r+1,['B','C']+[get_column_letter(4+i) for i in range(10)],['Floor plan','Units']+PQLAB)
r0=r+2
r=r0
for plan,u,vals in PRE:
    put(ws,f'B{r}',plan); put(ws,f'C{r}',u,BLUE,NUM)
    for i,v in enumerate(vals): put(ws,f'{get_column_letter(4+i)}{r}',v,BLUE,M0)
    r+=1
put(ws,f'B{r}','WEIGHTED AVG',BOLD)
put(ws,f'C{r}',f'=SUM(C{r0}:C{r-1})',BOLD,NUM)
for i in range(10):
    col=get_column_letter(4+i)
    put(ws,f'{col}{r}',f'=SUMPRODUCT({col}{r0}:{col}{r-1},$C${r0}:$C${r-1})/SUM($C${r0}:$C${r-1})',BOLD,M0)
PRE_W=r

# ---------------- L5_New_Leases ----------------
ws=sheet('L5_New_Leases',[2,22,14,14,14,50])
put(ws,'B2','Last-5 New Leases — Seasons model methodology',TITLE)
put(ws,'B3','Per floor plan: average of the 5 most recent new leases (by move-in) from the rent roll; property total = SUMPRODUCT(plan L5 avg x plan unit count) / total units — mirrors Seasons Assumptions L50:L58 / P49:T49.',NOTE,wrap=True)
put(ws,'B5','PRELUDE AT PARAMOUNT — RR w/ lease charges 8/13/2026 (current residents)',BOLD,fill=SUB)
l5=json.load(open('prelude_l5.json'))
mix={'pprA1':106,'pprB1':100,'pprB2':30,'pprC1':44}
r=6; plan_avg_rows={}
for p in ['pprA1','pprB1','pprB2','pprC1']:
    hdr(ws,r,['B','C','D'],[f'{p} — unit','Move-in','Rent'])
    r+=1; first=r
    for d,rent,uid in l5[p]:
        put(ws,f'B{r}',uid); put(ws,f'C{r}',datetime.datetime.strptime(d,'%Y-%m-%d'),BLUE,D); put(ws,f'D{r}',rent,BLUE,M0)
        r+=1
    put(ws,f'B{r}','L5 avg',BOLD); put(ws,f'D{r}',f'=AVERAGE(D{first}:D{r-1})',BOLD,M0)
    plan_avg_rows[p]=r
    r+=2
put(ws,f'B{r}','WEIGHTED SUMMARY',BOLD,fill=SUB)
hdr(ws,r+1,['B','C','D'],['Floor plan','Units (mix)','L5 avg'])
r0=r+2; r=r0
for p in ['pprA1','pprB1','pprB2','pprC1']:
    put(ws,f'B{r}',p); put(ws,f'C{r}',mix[p],BLUE,NUM); put(ws,f'D{r}',f'=D{plan_avg_rows[p]}',GREEN,M0)
    r+=1
put(ws,f'B{r}','PRELUDE L5 MIX-WEIGHTED',BOLD)
put(ws,f'D{r}',f'=SUMPRODUCT(D{r0}:D{r-1},C{r0}:C{r-1})/SUM(C{r0}:C{r-1})',BOLD,M0)
put(ws,f'F{r}','Result: $1,810.62 — vs Prelude UW Y1 market rent $1,730: +4.7%.',NOTE,wrap=True)
PL5=r
r+=2
put(ws,f'B{r}','SEASONS AT MERIDIAN — model Assumptions L50:L58 (8/4/26 RR), for comparison',BOLD,fill=SUB)
hdr(ws,r+1,['B','C','D'],['Floor plan','Units','L5 avg'])
sl5=[("S1_Seas",15,1495.8),("A1_Seas",120,1679),("A2_Seas",45,1766.4),("B1_Seas",30,1901.8),("B2_Seas",75,2074),("B3a_Seas",15,2012.4),("B3b_Seas",30,2066.2),("C1a_Seas",10,2497.8),("C1b_Seas",20,2266.4)]
r0=r+2; r=r0
for p,u,v in sl5:
    put(ws,f'B{r}',p); put(ws,f'C{r}',u,BLUE,NUM); put(ws,f'D{r}',v,BLUE,M0); r+=1
put(ws,f'B{r}','SEASONS L5 MIX-WEIGHTED',BOLD)
put(ws,f'D{r}',f'=SUMPRODUCT(D{r0}:D{r-1},C{r0}:C{r-1})/SUM(C{r0}:C{r-1})',BOLD,M0)
put(ws,f'F{r}','Ties to model Starting Market Rent $1,884.69.',NOTE)
SL5=r

# ---------------- Chart_Data ----------------
ws=sheet('Chart_Data',[2,13,15,15,15,15,15,17,15,44])
put(ws,'B2','Chart Data — plot-ready series',TITLE)
put(ws,'B3','One row per date. Green = computed on HD_Weighting / L5_New_Leases (mix-weighted); blue = extracted from models/statements. Quarterly HD points plotted at quarter midpoints.',NOTE,wrap=True)
hdr(ws,5,['B','C','D','E','F','G','H','I'],
    ['Date','Seasons HD asking (mix-wtd)','Seasons UW asking','Seasons L5 / UW entry','Prelude HD asking (mix-wtd)','Prelude UW rent/occ','Prelude actual in-place','Prelude L5 (8/13/26)'])
def qmid(lbl):
    q,y=lbl.split()[0],int(lbl.split()[1][:2])+2000
    m={'Q1':2,'Q2':5,'Q3':8,'Q4':11}[q]
    return datetime.datetime(y,m,15)
rows=[]
for i in range(3,13):  # Seasons real HD from Q2'24 (idx3) .. Q2'26 (idx11) + T90 idx12
    dt = qmid(QLAB[i]) if i<12 else datetime.datetime(2026,7,15)
    rows.append((dt,'C',f"=HD_Weighting!{get_column_letter(4+i)}{SEA_W}"))
for dt,v in [(datetime.datetime(2026,8,4),None)]:
    rows.append((dt,'E',f"=L5_New_Leases!D{SL5}"))
for dt,v in [(datetime.datetime(2026,8,4),1885),(datetime.datetime(2027,11,1),1960),(datetime.datetime(2028,11,1),2038),(datetime.datetime(2029,11,1),2110),(datetime.datetime(2030,11,1),2184),(datetime.datetime(2031,11,1),2249)]:
    rows.append((dt,'D',v))
for i in range(10):
    rows.append((qmid(PQLAB[i].replace(' QTD','')),'F',f"=HD_Weighting!{get_column_letter(4+i)}{PRE_W}"))
for dt,v in [(datetime.datetime(2026,7,1),1668),(datetime.datetime(2027,7,1),1745),(datetime.datetime(2028,7,1),1808),(datetime.datetime(2029,7,1),1883),(datetime.datetime(2030,7,1),1952)]:
    rows.append((dt,'G',v))
for m,v in enumerate([1668,1670,1672,1675,1675,1682,1684]):
    rows.append((datetime.datetime(2026,m+1,15),'H',v))
rows.append((datetime.datetime(2026,8,13),'I',f"=L5_New_Leases!D{PL5}"))
from collections import defaultdict
bydate=defaultdict(dict)
for dt,col,v in rows: bydate[dt][col]=v
r=6
for dt in sorted(bydate):
    put(ws,f'B{r}',dt,BLACK,D)
    for col,v in bydate[dt].items():
        if isinstance(v,str): put(ws,f'{col}{r}',v,GREEN,M0)
        else: put(ws,f'{col}{r}',v,BLUE,M0)
    r+=1
put(ws,f'B{r+1}','Sources: Seasons TMG model Rent Analysis (HD 8/2026) & Assumptions (UW path: $1,885 flat Y1, then +4/4/3.5/3.5/3% annual Nov steps); Prelude model Market Rent Summary (HD 11/2025) & Unleveraged (UW rent/occupied); Prelude accrual statement (2026 in-place = potential rent net LTL /280); rent roll 8/13/26 (L5). Seasons Q3 23-Q1 24 HD backfill excluded.',NOTE,wrap=True)
ws.merge_cells(f'B{r+1}:J{r+3}')

del wb['Sheet']
wb.move_sheet('Chart_Data', offset=-2)
wb.save('Rent_Chart_Data.xlsx')
print('saved')
