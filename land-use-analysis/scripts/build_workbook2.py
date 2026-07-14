"""Enhanced, granular land-use workbook (supersedes build_workbook.py for this study).
Adds vs the base builder:
  * Future Land Use (Comp-Plan) sheet + FLU designation/intent on every parcel
  * Parcels sheet enriched with subdivision, assessor value, owner-occupied (homestead), FLU
  * Developable Inventory sheet — the granular per-parcel developable set (no owner mega-blobs),
    with contiguous-block grouping + a research-priority score
  * Ownership & Intent sheet — the researched dossiers (owner / since-when / who / intent)
  * Supply Summary sheet — area rollups (the comparison backbone)
Reads: config.json, in/parcels_enriched.geojson, in/developable.json, Tables/supply_summary.csv,
Tables/flu_summary.csv, in/dossiers.json (optional), Tables/decisions_log.md (optional).
"""
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config as C  # noqa: E402
import palette as P  # noqa: E402
from enrich import FLU_INTENT  # noqa: E402

SERIF = "Cambria"
THREAT_FILL = {"High": "F4C7C3", "Medium": "FCE3C6", "Low": "FFF7D6", "Unknown": "E4D7F0"}
STATUS_FILL = {"approved": "C6E0B4", "under_construction": "BDD7EE", "proposed": "FFE699",
               "denied": "F4B7B7", "built": "D9D9D9", "dormant": "F2F2F2", "none": "F2F2F2"}


def _csv(path):
    return list(csv.DictReader(open(path, encoding="utf-8"))) if Path(path).exists() else []


def threat_word(s):
    return (str(s or "").split("-")[0].split()[0].strip().capitalize() if s else "")


def main():
    cfg = C.load()
    TBL, IN = C.tables(cfg), C.indir(cfg)
    brand = cfg.get("branding", {})
    NAVY = brand.get("primary", "#0B4F6C").lstrip("#")
    GRAY, INK, BAND = "6F6F70", "1A2B3C", "F1F6FA"
    thin = Side(style="thin", color="D9D9D9")
    B = Border(thin, thin, thin, thin)
    name = cfg["subject"]["name"]

    wb = Workbook()
    wb.remove(wb.active)

    def title(ws, text, sub, ncols):
        ws.sheet_view.showGridLines = False
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        c = ws.cell(1, 1, text); c.font = Font(name=SERIF, size=18, bold=True, color=NAVY)
        ws.row_dimensions[1].height = 28
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        ws.cell(2, 1, sub).font = Font(name="Calibri", size=10, italic=True, color=GRAY)

    def header(ws, row, cols):
        for j, h in enumerate(cols, 1):
            c = ws.cell(row, j, h)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            c.alignment = Alignment("center", "center", wrap_text=True); c.border = B
        ws.row_dimensions[row].height = 30

    def body(ws, start, rows, widths, wrapcols=(), threatcol=None, statuscol=None,
             catcolorcol=None, catcolors=None):
        for i, r in enumerate(rows):
            rr = start + i
            for j, v in enumerate(r, 1):
                c = ws.cell(rr, j, v); c.border = B
                c.font = Font(name="Calibri", size=10, color=INK)
                c.alignment = Alignment(vertical="top", wrap_text=(j - 1 in wrapcols))
                if i % 2:
                    c.fill = PatternFill("solid", fgColor=BAND)
                if threatcol is not None and j - 1 == threatcol:
                    tw = threat_word(v)
                    if tw in THREAT_FILL:
                        c.fill = PatternFill("solid", fgColor=THREAT_FILL[tw])
                        c.font = Font(name="Calibri", size=10, bold=True, color=INK)
                if statuscol is not None and j - 1 == statuscol:
                    key = str(v or "").split("/")[0].strip().lower().replace(" ", "_")
                    if key in STATUS_FILL:
                        c.fill = PatternFill("solid", fgColor=STATUS_FILL[key])
                if catcolorcol is not None and j - 1 == catcolorcol and catcolors and catcolors[i]:
                    c.fill = PatternFill("solid", fgColor=catcolors[i].lstrip("#"))
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(start, 1)
        return start + len(rows)

    def paragraphs(ws, start, text, width=112):
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = width
        r = start
        for line in text.splitlines():
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            cell = ws.cell(r, 1)
            bold = line.startswith("#") or re.match(r"^\s*\d+\.\s", line) or line.startswith("**")
            txt = re.sub(r"^#+\s*", "", line).replace("**", "")
            cell.value = txt
            cell.font = Font(name="Calibri", size=11 if line.startswith("# ") else 10,
                             bold=bool(bold), color=NAVY if line.startswith("#") else INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[r].height = max(15, 15 * (1 + len(txt) // width))
            r += 1
        return r

    def lean_sheet(sheetname, title_text, sub, cols, rows, widths, threatcol=None):
        ws = wb.create_sheet(sheetname)
        title(ws, title_text, sub, len(cols))
        header(ws, 4, cols)
        for r in rows:
            ws.append(list(r))
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        last = 4 + len(rows)
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(last, 4)}"
        if threatcol is not None and rows:
            col = get_column_letter(threatcol + 1)
            for tw, fill in THREAT_FILL.items():
                ws.conditional_formatting.add(f"{col}5:{col}{last}", CellIsRule(
                    operator="equal", formula=[f'"{tw}"'], fill=PatternFill("solid", fgColor=fill)))
        return ws

    # load data
    parcels = json.loads((IN / "parcels_enriched.geojson").read_text())["features"]
    dev = json.loads((IN / "developable.json").read_text()) if (IN / "developable.json").exists() else []
    dev_accts = {r["account"] for r in dev}
    supply = _csv(TBL / "supply_summary.csv")
    dossiers = json.loads((IN / "dossiers.json").read_text()) if (IN / "dossiers.json").exists() else []
    a = cfg["analysis_area"]
    area_desc = f"{a.get('radius_mi')}-mile radius"

    # ── Overview ─────────────────────────────────────────────────────────────
    ws = wb.create_sheet("Overview")
    title(ws, f"{name} — Land Use & Developable-Land Analysis",
          "Who owns the developable land nearby, since when, who they are, and what they plan", 2)
    sub = cfg["subject"]
    total = next((r for r in supply if r["Metric"].startswith("TOTAL developable")), {})
    highrow = next((r for r in supply if r["Group"] == "High"), {})
    info = [
        ("Subject", name), ("Address", sub.get("address")),
        ("Verified point", f"{sub.get('lat')}, {sub.get('lon')}"),
        ("Location note", sub.get("location_note")),
        ("Analysis area", area_desc + f" ({len(parcels):,} parcels classified)"),
        ("Developable vacant land", f"{total.get('Parcels','')} parcels · {total.get('Acres','')} acres"),
        ("  of which MF by-right (High)", f"{highrow.get('Parcels','')} parcels · {highrow.get('Acres','')} acres"),
        ("Vacant definition", cfg.get("landuse_scheme", {}).get("vacant_note")),
        ("Min. developable acreage", cfg.get("min_acres", 1.0)),
        ("Owner data note", "Ada County's public GIS/portal does not publish owner names (Idaho); "
         "ownership, tenure, and intent for the material sites are researched per-site (see "
         "Ownership & Intent) from development applications, city staff reports, and local reporting."),
    ]
    r = 4
    for k, v in info:
        ws.cell(r, 1, k).font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        c = ws.cell(r, 2, v); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="Calibri", size=10, color=INK); r += 1
    hl = (TBL / "headline.txt").read_text(encoding="utf-8") if (TBL / "headline.txt").exists() else \
        cfg.get("headline", "(see comparison report)")
    ws.cell(r + 1, 1, "Headline").font = Font(name=SERIF, size=12, bold=True, color=NAVY)
    ws.merge_cells(start_row=r + 2, start_column=1, end_row=r + 8, end_column=2)
    hc = ws.cell(r + 2, 1, hl); hc.alignment = Alignment(wrap_text=True, vertical="top")
    hc.font = Font(name="Calibri", size=10, color=INK)
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 100

    # ── Assumptions & Decisions ──────────────────────────────────────────────
    ws = wb.create_sheet("Assumptions & Decisions")
    dlog = TBL / "decisions_log.md"
    title(ws, "Assumptions & Decisions", "Every subjective judgment call + data note", 6)
    if dlog.exists():
        paragraphs(ws, 4, dlog.read_text(encoding="utf-8"))
    else:
        ws.cell(4, 1, "decisions_log.md not found.")

    # ── Ownership & Intent (dossiers) ────────────────────────────────────────
    ws = wb.create_sheet("Ownership & Intent")
    title(ws, "Ownership & Development Intent — Key Developable Sites",
          "Who owns each material site, since when, who they are, and what they plan (researched, sourced)", 11)
    header(ws, 4, ["Site", "Location", "Acres", "Zoning / FLU", "Owner (as researched)", "Owner type",
                   "Owned since", "Who they are", "Intent — what they'll do", "Status", "Sources"])
    drows, dcolors = [], []
    for d in dossiers:
        drows.append([d.get("site_id"), d.get("address"), d.get("acres"),
                      d.get("zoning_flu", d.get("zoning", "")), d.get("owner_name"),
                      d.get("owner_type"), d.get("ownership_since"), d.get("who_they_are"),
                      d.get("intent_summary"), d.get("project_status"),
                      "; ".join(d.get("sources", [])) if isinstance(d.get("sources"), list) else d.get("sources")])
    body(ws, 5, drows, [8, 22, 7, 22, 24, 18, 14, 34, 46, 14, 40],
         wrapcols=(1, 3, 4, 7, 8, 10), statuscol=9)

    # ── Supply Summary ───────────────────────────────────────────────────────
    ws = wb.create_sheet("Supply Summary")
    title(ws, "Developable-Land Supply Summary", "Area rollups — the comparison backbone", 4)
    header(ws, 4, ["Metric", "Group", "Parcels", "Acres"])
    body(ws, 5, [[r["Metric"], r["Group"], r["Parcels"], r["Acres"]] for r in supply],
         [34, 30, 12, 12], wrapcols=(0,))

    # ── Developable Inventory ────────────────────────────────────────────────
    dv_rows = [[r["account"], r["situs"], r["subdivision"], r["acres"], r["contig_group"],
                r["contig_group_acres"], r["jurisdiction"], r["zone_plain"], r["mf_threat"],
                r["flu_plain"], r["flu_intent"], r["total_value"], r["dist_mi"], r["interest"]]
               for r in dev]
    lean_sheet("Developable Inventory",
               f"Developable Vacant Land — Per Parcel ({len(dev):,})",
               "Granular per-parcel developable inventory (HOA-common/floodplain/sliver-filtered; "
               "contiguous blocks grouped, NOT fused). Sorted by research-priority score.",
               ["Account", "Address", "Subdivision", "Acres", "Contig Grp", "Grp Acres",
                "Jurisdiction", "Zoning", "MF Threat", "Future Land Use", "FLU Intent",
                "Assessor Value", "Dist (mi)", "Priority"],
               dv_rows, [16, 24, 22, 8, 9, 9, 16, 26, 10, 26, 22, 13, 8, 8], threatcol=8)

    # ── Parcels (enriched) ───────────────────────────────────────────────────
    prows = []
    for f in parcels:
        p = f["properties"]
        prows.append([p.get("account"), p.get("situs"), p.get("subdivision"),
                      p.get("landuse_code"), p.get("bucket"), p.get("jurisdiction"),
                      p.get("zone_plain"), p.get("zone_category"), p.get("mf_threat"),
                      p.get("flu_plain"), p.get("flu_intent"), p.get("total_value"),
                      "Yes" if p.get("owner_occupied") else "", "Yes" if p.get("is_vacant") else "",
                      "Yes" if p.get("account") in dev_accts else "",
                      "Yes" if p.get("is_hoa_common") else "", p.get("acres"), p.get("dist_mi")])
    prows.sort(key=lambda r: (r[17] if isinstance(r[17], (int, float)) else 9e9))
    lean_sheet("Parcels", f"All Parcels in Area ({len(prows):,})",
               "Every parcel: land use, zoning, MF threat, Future Land Use, value, owner-occupancy — sorted by distance",
               ["Account", "Address", "Subdivision", "Use", "Land-Use Bucket", "Jurisdiction",
                "Zoning", "Zone Category", "MF Threat", "Future Land Use", "FLU Intent",
                "Assessor Value", "Owner-Occ", "Vacant?", "Developable?", "HOA/Common?",
                "Acres", "Dist (mi)"],
               prows, [16, 24, 20, 6, 22, 16, 24, 22, 10, 24, 20, 12, 9, 8, 11, 11, 8, 8], threatcol=8)

    # ── Land-Use Grouping ────────────────────────────────────────────────────
    rows_lu = _csv(TBL / "land_use_buckets.csv")
    ws = wb.create_sheet("Land-Use Grouping")
    title(ws, "Land-Use Grouping", "How each assessor property code becomes a land-use category", 5)
    header(ws, 4, ["Land-Use Category", "Use Code", "What the code means", "Parcels in area", "Source"])
    lr = [[r["Bucket"], r["Code"], r["Description"], r.get("Parcels_in_area", ""), r["Source"]] for r in rows_lu]
    lc = [P.landuse_color(r["Bucket"]) for r in rows_lu]
    body(ws, 5, lr, [30, 12, 52, 14, 40], wrapcols=(2, 4), catcolorcol=0, catcolors=lc)

    # ── Zoning Grouping ──────────────────────────────────────────────────────
    rows_z = _csv(TBL / "zoning_categories.csv")
    zplain = cfg.get("zone_plain", {})
    ws = wb.create_sheet("Zoning Grouping")
    title(ws, "Zoning Grouping (7 Ada County jurisdictions)",
          "How each district is grouped and its multifamily-supply threat if vacant", 6)
    header(ws, 4, ["Jurisdiction", "Zone", "What the district is", "Grouped Category",
                   "MF Threat", "What the threat means"])
    torder = {"High": 0, "Medium": 1, "Unknown": 2, "Low": 3}
    rows_z.sort(key=lambda r: (r["Jurisdiction"], torder.get(r["MF Threat (if vacant)"], 9), r["Base Zone"]))
    zr = [[r["Jurisdiction"], r["Base Zone"], zplain.get(r["Base Zone"], r["Base Zone"]),
           r["Category"], r["MF Threat (if vacant)"],
           P.THREAT_DEF.get(threat_word(r["MF Threat (if vacant)"]), "")] for r in rows_z]
    body(ws, 5, zr, [20, 8, 30, 30, 12, 50], wrapcols=(2, 3, 5), threatcol=4)

    # ── Future Land Use ──────────────────────────────────────────────────────
    flu = _csv(TBL / "flu_summary.csv")
    ws = wb.create_sheet("Future Land Use")
    title(ws, "Future Land Use (Comprehensive Plan)",
          "Adopted policy designation = what the land is PLANNED for (a rezone-likely signal for vacant land)", 4)
    header(ws, 4, ["FLU designation (source: code)", "Plain meaning", "Residential-intensification intent",
                   "Vacant parcels in area"])
    frows = []
    for r in flu:
        raw = r["FLU (source: designation)"]
        desig = raw.split(":", 1)[1].strip() if ":" in raw else raw
        plain, intent = FLU_INTENT.get(desig, (desig, ""))
        frows.append([raw, plain, intent, r["Vacant parcels in area"]])
    body(ws, 5, frows, [40, 34, 30, 16], wrapcols=(1, 2))

    # ── Data Sources ─────────────────────────────────────────────────────────
    ws = wb.create_sheet("Data Sources")
    title(ws, "Data Sources", "Every public layer used, with its endpoint", 4)
    header(ws, 4, ["Type", "Jurisdiction / Name", "Endpoint", "Notes"])
    srows = [["Parcels", s.get("name"), s["url"], f"confidence: {s.get('data_confidence','full')}; "
              "no owner field (Idaho)"] for s in cfg.get("parcel_sources", [])]
    zj = sorted({s["jurisdiction"] for s in cfg.get("zoning_sources", [])})
    srows.append(["Zoning", ", ".join(zj), cfg["zoning_sources"][0]["url"],
                  "one countywide layer, split by CITY field into 7 jurisdictions"])
    srows.append(["Future Land Use", "Boise / Meridian / Ada County", "Ada County GIS FLU FeatureServers",
                  "adopted comprehensive-plan designations"])
    srows.append(["Ownership & Intent", "development applications, city staff reports, BoiseDev, COMPASS, Idaho SOS",
                  "(see Ownership & Intent sheet sources)", "researched per site"])
    body(ws, 5, srows, [14, 30, 74, 34], wrapcols=(1, 2, 3))

    # ── Methodology ──────────────────────────────────────────────────────────
    ws = wb.create_sheet("Methodology")
    title(ws, "Methodology", "How the analysis was produced", 6)
    paragraphs(ws, 4, (TBL / "methodology.md").read_text(encoding="utf-8")
               if (TBL / "methodology.md").exists() else "See Assumptions & Decisions.")

    order = ["Overview", "Assumptions & Decisions", "Ownership & Intent", "Supply Summary",
             "Developable Inventory", "Parcels", "Land-Use Grouping", "Zoning Grouping",
             "Future Land Use", "Data Sources", "Methodology"]
    wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)

    safe = "".join(ch if ch.isalnum() or ch in " -_" else "" for ch in name).strip()
    out = C.root(cfg) / f"{safe} - Land Use Analysis.xlsx"
    wb.save(out)
    print(f"wrote {out.name}  ({len(prows):,} parcels, {len(dev):,} developable, {len(dossiers)} dossiers)")


if __name__ == "__main__":
    main()
