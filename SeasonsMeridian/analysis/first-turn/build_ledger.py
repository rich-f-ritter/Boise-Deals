#!/usr/bin/env python3
"""
Seasons at Meridian — First-Turn Analysis: ledger assembly.

Builds a per-tenant ledger for the 1/1/2026 resident cohort, tracking each
initial tenant to the 8/4/2026 rent roll: renewed / moved out / MTM / not yet
expired, with gross and effective rents on both sides of the turn.

Sources (see ../../documents/ and DATA_NOTES.md for caveats):
  - Rent rolls: 1/1/2026 (BACKDATED: roster+rents reliable, dates are current),
    7/07, 7/19, 8/04/2026 (contemporaneous)
  - Concession Burn Off 6/21 & 7/30/2026 (only source of Lease Start dates)
  - Renewal Tradeouts report 5/10-7/9/2026 (exact prior/new concession detail)
  - HelloData unit details (executed new-lease asking/effective rents)

Output: ledger.csv + validation report to stdout.
"""
import csv
import openpyxl
from datetime import datetime, date, timedelta
from collections import Counter

DOCS = '../../documents'
CUTOFF = date(2026, 8, 4)          # measurement date = latest rent roll
RENEWAL_MIN_GAP_DAYS = 45          # lease start > move-in + gap => renewal
CORPORATE_RESIDENTS = {'Coleman Environmental Engineering'}


def d(x):
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    if isinstance(x, str):
        s = x.replace('*', '').strip()
        for fmt in ('%m/%d/%Y', '%Y-%m-%d'):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    return None


def fnum(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def parse_rent_roll(fname):
    wb = openpyxl.load_workbook(fname, read_only=True, data_only=True)
    ws = wb['Report1']
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    section = None
    recs = []
    for r in rows[6:]:
        c0 = r[0]
        if isinstance(c0, str) and ('Current/Notice' in c0 or 'Future' in c0):
            section = c0.strip()
        if c0 is None or r[1] is None:
            continue
        recs.append({
            'section': section, 'unit': str(c0).strip(), 'type': str(r[1]).strip(),
            'sf': fnum(r[2]), 'res': str(r[3]).strip() if r[3] else '',
            'name': str(r[4]).strip() if r[4] else '',
            'mkt': fnum(r[5]), 'actual': fnum(r[6]),
            'movein': d(r[9]), 'leaseexp': d(r[10]), 'moveout': d(r[11]),
        })
    return [r for r in recs if r['section'] and 'Current' in r['section']]


def parse_burnoff(fname):
    """Section 1 only (the burn-off table; section 2 'Projection by Unit' is skipped)."""
    wb = openpyxl.load_workbook(fname, read_only=True, data_only=True)
    ws = wb['Report1']
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    recs = {}
    for r in rows[6:]:
        c0 = r[0]
        if isinstance(c0, str) and 'Projection by Unit' in c0:
            break
        if c0 is None or r[2] is None or not str(r[2]).startswith('t'):
            continue
        recs[str(r[2]).strip()] = {
            'unit': str(c0).strip(), 'movein': d(r[4]), 'leasestart': d(r[5]),
            'tot_conc': fnum(r[6]), 'cur_conc': fnum(r[7]),
            'term': fnum(r[10]), 'mkt': fnum(r[11]), 'leaserent': fnum(r[12]),
        }
    return recs


def parse_renewal_report(fname):
    wb = openpyxl.load_workbook(fname, read_only=True, data_only=True)
    ws = wb['(3pseason) Seasons at Meridian']
    recs = {}
    for row in ws.iter_rows(min_row=8, values_only=True):
        if row[4] is None or str(row[4]).strip() == '':
            continue
        recs[str(row[4]).strip()] = {
            'name': row[6], 'old_exp': d(row[7]), 'new_start': d(row[8]),
            'prior_term': fnum(row[9]), 'prior_gross': fnum(row[10]),
            'prior_conc': fnum(row[11]), 'prior_eff': fnum(row[13]),
            'new_term': fnum(row[14]), 'new_gross': fnum(row[15]),
            'new_conc': fnum(row[16]), 'new_eff': fnum(row[18]),
        }
    wb.close()
    return recs


def parse_hellodata(fname):
    with open(fname) as f:
        hd = list(csv.DictReader(f))
    by_unit = {}
    for r in hd:
        r['off'] = d(r['Off Market Date'])
        r['ask'] = fnum(r['Last Asking Rent'])
        r['eff'] = fnum(r['Last Effective Rent'])
        r['term_hd'] = fnum(r['Term'])
        by_unit.setdefault(r['Unit'], []).append(r)
    return by_unit


def hd_episode(hd_by_unit, unit, movein):
    """Best HelloData episode for a move-in: off-market within -120..+30 days."""
    best = None
    for c in hd_by_unit.get(unit, []):
        if c['off'] is None or movein is None:
            continue
        lag = (movein - c['off']).days           # positive = off-market before move-in
        if -30 <= lag <= 120:
            if best is None or abs(lag - 15) < abs((movein - best['off']).days - 15):
                best = c
    return best


def months_between(a, b):
    """Approximate months from a to b."""
    return round((b - a).days / 30.44) if a and b else None


def main():
    rr_jan = parse_rent_roll(f'{DOCS}/rent-rolls/RentRoll_AsOf_2026-01-01_BACKDATED-see-notes.xlsx')
    rr_aug = parse_rent_roll(f'{DOCS}/rent-rolls/RentRoll_AsOf_2026-08-04.xlsx')
    rr_jul = parse_rent_roll(f'{DOCS}/rent-rolls/RentRoll_AsOf_2026-07-07.xlsx')
    bo_jun = parse_burnoff(f'{DOCS}/concession-burnoff/ConcessionBurnOff_AsOf_2026-06-21.xlsx')
    bo_jul = parse_burnoff(f'{DOCS}/concession-burnoff/ConcessionBurnOff_AsOf_2026-07-30.xlsx')
    renew = parse_renewal_report(f'{DOCS}/renewal-reports/RenewalTradeouts_2026-05-10_to_2026-07-09.xlsx')
    hd = parse_hellodata(f'{DOCS}/hellodata/HelloData_UnitDetails_2026-08-12.csv')

    cohort = [r for r in rr_jan if r['res'].startswith('t')]
    aug_by_res = {r['res']: r for r in rr_aug if r['res'].startswith('t')}
    jul_by_res = {r['res']: r for r in rr_jul if r['res'].startswith('t')}
    aug_by_unit = {}
    for r in rr_aug:
        if r['res'].startswith('t'):
            aug_by_unit.setdefault(r['unit'], []).append(r)

    ledger = []
    val = Counter()

    for c in cohort:
        res, unit = c['res'], c['unit']
        aug = aug_by_res.get(res)
        bo = bo_jul.get(res) or bo_jun.get(res)
        rr = renew.get(unit)
        if rr and aug is None:
            rr = None                                   # renewal report row belongs to someone else
        if rr and rr.get('name') and aug and rr['name'] != aug['name']:
            rr = None
        hd_init = hd_episode(hd, unit, c['movein'])

        row = {
            'unit': unit, 'unit_type': c['type'], 'sf': c['sf'],
            'resident_id': res, 'name': c['name'], 'move_in': c['movein'],
            'exclude': 'Y' if c['name'] in CORPORATE_RESIDENTS else '',
        }

        # ---- initial lease: gross rent ----
        # Jan-roll Actual Rent = charge as of Jan 2026 (validated vs renewal report priors)
        init_gross, init_gross_src = c['actual'], 'RR 1/1/26'
        if not init_gross:                              # Yardi zero-charge quirk on a few rows
            bo6 = bo_jun.get(res)
            if bo6 and bo6['leaserent'] and bo6['leasestart'] and bo6['movein'] and \
                    (bo6['leasestart'] - bo6['movein']).days <= RENEWAL_MIN_GAP_DAYS:
                init_gross, init_gross_src = bo6['leaserent'], 'Burn-off 6/21 (Jan roll rent was 0)'
            elif hd_init and hd_init['ask']:
                init_gross, init_gross_src = hd_init['ask'], 'HelloData episode (Jan roll rent was 0)'
            else:
                init_gross, init_gross_src = None, 'UNKNOWN (Jan roll rent was 0)'
            val['Jan roll rent 0 -> fallback'] += 1
        # ---- renewal detection ----
        renewed = False
        renewal_date = None
        if aug and bo and bo['leasestart'] and bo['movein'] and \
                (bo['leasestart'] - bo['movein']).days > RENEWAL_MIN_GAP_DAYS:
            renewed = True
            renewal_date = bo['leasestart']
        if aug and rr and not renewed:                  # in report but missed by burn-off
            renewed = True
            renewal_date = rr['new_start']
            val['renewal from report only'] += 1

        # ---- original (first-lease) expiration ----
        if renewed:
            if rr and rr['old_exp']:
                orig_exp, orig_exp_src = rr['old_exp'], 'Renewal report'
            else:
                orig_exp, orig_exp_src = renewal_date - timedelta(days=1), 'Renewal lease start - 1d'
            if renewal_date < date(2026, 1, 1):
                # renewed before the Jan snapshot: Jan rent is already the renewal rent;
                # recover the initial gross from HelloData
                init_gross_src = 'HelloData episode (pre-Jan renewal)'
                init_gross = hd_init['ask'] if hd_init else None
                val['pre-Jan renewals'] += 1
        else:
            src_roll = aug or jul_by_res.get(res) or c
            orig_exp, orig_exp_src = src_roll['leaseexp'], 'Rent roll (still on initial lease)'

        # ---- initial lease: term + concessions -> effective ----
        init_term = None
        init_conc, init_conc_src = None, ''
        if rr:
            init_term = rr['prior_term']
            init_conc, init_conc_src = rr['prior_conc'], 'Renewal report'
        elif not renewed and bo:
            init_term = bo['term']
            init_conc = -bo['cur_conc'] if bo['cur_conc'] is not None else None
            init_conc_src = 'Burn-off'
        elif hd_init:
            init_term = hd_init['term_hd']
            if hd_init['ask'] and hd_init['eff'] is not None and hd_init['term_hd']:
                init_conc = (hd_init['ask'] - hd_init['eff']) * hd_init['term_hd']
                init_conc_src = 'HelloData (ask-eff x term)'
        if init_term is None:
            init_term = months_between(c['movein'], orig_exp)
            if init_term is not None:
                val['term inferred from dates'] += 1
        init_eff = None
        if init_gross is not None and init_term:
            init_eff = init_gross - (init_conc or 0) / init_term
            if init_conc is None:
                init_conc_src = 'None found (eff = gross)'
                val['initial concession unknown -> assumed 0'] += 1

        # ---- outcome classification ----
        expired_by_cutoff = orig_exp is not None and orig_exp <= CUTOFF
        renewal_gross = renewal_term = renewal_conc = renewal_eff = None
        vacated = new_res = new_name = None
        new_movein = new_gross = new_term = new_conc = new_eff = None
        days_vacant = None

        if aug:
            if renewed:
                outcome = 'Renewed' if expired_by_cutoff else 'Renewed Early (exp not yet due)'
                renewal_gross = aug['actual']
                renewal_term = (rr and rr['new_term']) or (bo and bo['term'])
                renewal_conc = rr['new_conc'] if rr else (-bo['cur_conc'] if bo and bo['cur_conc'] is not None else 0)
                if renewal_gross is not None and renewal_term:
                    renewal_eff = renewal_gross - (renewal_conc or 0) / renewal_term
            elif expired_by_cutoff:
                outcome = 'MTM Holdover'
            else:
                outcome = 'Not Yet Expired'
            if aug['moveout']:
                row['notice'] = 'Y'
        else:
            outcome = 'Moved Out'
            src = jul_by_res.get(res) or c
            vacated = src['moveout']
            early = vacated and orig_exp and (orig_exp - vacated).days > 5
            row['early_term'] = 'Y' if (early or not expired_by_cutoff) else ''
            repl = [x for x in aug_by_unit.get(unit, []) if x['res'] != res]
            if repl:
                n = repl[0]
                new_res, new_name, new_movein = n['res'], n['name'], n['movein']
                new_gross = n['actual']
                nbo = bo_jul.get(n['res']) or bo_jun.get(n['res'])
                hd_new = hd_episode(hd, unit, n['movein'])
                new_term = (nbo and nbo['term']) or (hd_new and hd_new['term_hd'])
                if nbo and nbo['cur_conc'] is not None:
                    new_conc = -nbo['cur_conc']
                elif hd_new and hd_new['ask'] and hd_new['eff'] is not None and new_term:
                    new_conc = (hd_new['ask'] - hd_new['eff']) * new_term
                if new_gross is not None and new_term:
                    new_eff = new_gross - (new_conc or 0) / new_term
                if vacated and new_movein:
                    days_vacant = (new_movein - vacated).days

        in_denominator = expired_by_cutoff and outcome in ('Renewed', 'MTM Holdover', 'Moved Out')
        row.update({
            'initial_term_mo': init_term, 'orig_expiration': orig_exp,
            'orig_exp_source': orig_exp_src,
            'initial_gross': init_gross, 'initial_gross_source': init_gross_src,
            'initial_conc_total': round(init_conc, 2) if init_conc is not None else None,
            'initial_conc_source': init_conc_src,
            'initial_eff': round(init_eff, 2) if init_eff is not None else None,
            'outcome': outcome,
            'expired_by_cutoff': 'Y' if expired_by_cutoff else '',
            'in_retention_denom': 'Y' if in_denominator else '',
            'renewal_date': renewal_date,
            'renewal_term_mo': renewal_term,
            'renewal_gross': renewal_gross,
            'renewal_conc_total': round(renewal_conc, 2) if renewal_conc is not None else None,
            'renewal_eff': round(renewal_eff, 2) if renewal_eff is not None else None,
            'renewal_in_report': 'Y' if rr else '',
            'vacated_date': vacated,
            'new_resident_id': new_res, 'new_name': new_name, 'new_move_in': new_movein,
            'new_lease_term_mo': new_term, 'new_gross': new_gross,
            'new_conc_total': round(new_conc, 2) if new_conc is not None else None,
            'new_eff': round(new_eff, 2) if new_eff is not None else None,
            'days_vacant': days_vacant,
        })
        row.setdefault('notice', '')
        row.setdefault('early_term', '')
        ledger.append(row)
        val[outcome] += 1

    # ---------- validation ----------
    print('=== outcome counts ===')
    for k, v in val.most_common():
        print(f'  {v:4}  {k}')
    denom = [r for r in ledger if r['in_retention_denom'] == 'Y']
    ren = [r for r in denom if r['outcome'] == 'Renewed']
    mtm = [r for r in denom if r['outcome'] == 'MTM Holdover']
    out = [r for r in denom if r['outcome'] == 'Moved Out']
    print(f'\nretention: {len(ren)} renewed / {len(denom)} expirations = {len(ren)/len(denom):.1%}'
          f'   (moved out {len(out)}, MTM {len(mtm)})')
    early = [r for r in ledger if r['outcome'] == 'Moved Out' and r['early_term'] == 'Y']
    print(f'early terminations (not in denominator unless exp since passed): {len(early)}')

    # renewal report reconciliation: every in-window renewal must be flagged Renewed
    missed = [u for u in renew if not any(
        r['unit'] == u and r['outcome'].startswith('Renewed') for r in ledger)]
    print(f'renewal-report units not classified Renewed in ledger: {missed}')

    # trade-out previews (exclusions applied)
    def stats(pairs):
        pct = [(b - a) / a * 100 for a, b in pairs if a and b]
        pct.sort()
        n = len(pct)
        return f'n={n} avg {sum(pct)/n:+.1f}% median {pct[n//2]:+.1f}%' if n else 'n=0'
    ren_ok = [r for r in ren if not r['exclude']]
    print('renewal gross trade-out:',
          stats([(r['initial_gross'], r['renewal_gross']) for r in ren_ok]))
    print('renewal eff trade-out:  ',
          stats([(r['initial_eff'], r['renewal_eff']) for r in ren_ok]))
    turns = [r for r in ledger if r['outcome'] == 'Moved Out' and r['new_gross'] and not r['exclude']]
    print('new-lease gross trade-out:',
          stats([(r['initial_gross'], r['new_gross']) for r in turns]))
    print('new-lease eff trade-out:  ',
          stats([(r['initial_eff'], r['new_eff']) for r in turns]))

    cols = list(ledger[0].keys())
    with open('ledger.csv', 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(ledger)
    print(f'\nwrote ledger.csv ({len(ledger)} rows)')


if __name__ == '__main__':
    main()
